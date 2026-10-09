"""Build the 15-page English report as Times New Roman 13pt Word and PDF."""
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader
from common import ROOT, write_json
from report_content import PAGES

FONT='Times New Roman'
MEMBERS=[('Lã Minh Trung','20251319M'),('Nguyễn Thu Uyên','20252279M'),('Nguyễn Thị Nhã Linh','20261262M'),('Nguyễn Khắc Thái Bình','20251324M')]
SUPERVISOR='TS. Đỗ Bá Lâm'

def contents(doc):
    """Standard paragraph-based contents with right-aligned page numbers."""
    for number,page in enumerate(PAGES[2:],3):
        prefix,label=page['title'].split('. ',1)
        text=prefix+'. '+label.capitalize()
        # Preserve the standard capitalization of technical acronyms.
        for term in ['RDF','OWL','SPARQL']:text=text.replace(term.lower(),term)
        entry=paragraph(doc,text+'\t'+str(number),bold=True,align=WD_ALIGN_PARAGRAPH.LEFT,first_line=False)
        entry.paragraph_format.space_after=Pt(4)
        entry.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        for block in page['blocks']:
            if block['type']!='heading':continue
            entry=paragraph(doc,block['text']+'\t'+str(number),align=WD_ALIGN_PARAGRAPH.LEFT,first_line=False)
            entry.paragraph_format.left_indent=Cm(.75);entry.paragraph_format.space_after=Pt(4)
            entry.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)

def format_run(run, bold=False, italic=False):
    run.font.name=FONT;run.font.size=Pt(13);run.font.bold=bold;run.font.italic=italic
    run.font.color.rgb=RGBColor(0,0,0)
    fonts=run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+name),FONT)

def paragraph(doc,text='',bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,first_line=True):
    p=doc.add_paragraph();p.alignment=align
    f=p.paragraph_format;f.line_spacing=1.5;f.space_before=Pt(0);f.space_after=Pt(6)
    f.first_line_indent=Cm(1) if first_line else Cm(0)
    f.widow_control=True
    format_run(p.add_run(text),bold,italic)
    return p

def heading(doc,text):
    p=paragraph(doc,text,bold=True,align=WD_ALIGN_PARAGRAPH.LEFT,first_line=False)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(5)
    return p

def figure(doc,block):
    path=ROOT/block['path'];w,h=Image.open(path).size
    width=min(block['width'],block['max_height']*w/h)
    p=paragraph(doc,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    picture=p.add_run().add_picture(str(path),width=Cm(width))
    picture._inline.docPr.set('descr',block['caption'])
    caption=paragraph(doc,block['caption'],italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    caption.paragraph_format.keep_together=True

def add_table(doc,block):
    if block.get('caption'):
        p=paragraph(doc,block['caption'],italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
        p.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1, cols=len(block['headers']));table.style='Table Grid';table.autofit=False
    widths=block.get('widths') or [16/len(block['headers'])]*len(block['headers'])
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    for i,values in enumerate([block['headers']]+block['rows']):
        row=table.rows[0] if i==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        for cell,width,value in zip(row.cells,widths,values):
            cell.width=Cm(width);p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.first_line_indent=Cm(0)
            format_run(p.add_run(str(value)),bold=i==0)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E6E6E6');cell._tc.get_or_add_tcPr().append(shading)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    after=paragraph(doc,first_line=False);after.paragraph_format.space_after=Pt(0);after.paragraph_format.line_spacing=1.5
    # A structural spacer carries no visible text; minimize its occupied space.
    after.paragraph_format.line_spacing=Pt(2)

def cover(doc):
    for text in ['HANOI UNIVERSITY OF SCIENCE AND TECHNOLOGY','SEMANTIC WEB COURSE']:
        paragraph(doc,text,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    logo=paragraph(doc,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    logo.paragraph_format.line_spacing=1
    picture=logo.add_run().add_picture(str(ROOT/'docs/report_images/hust_logo.png'),height=Cm(4))
    picture._inline.docPr.set('descr','Official Hanoi University of Science and Technology logo')
    paragraph(doc,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    for text in ['COURSE PROJECT REPORT','MOVIELOD','A LINKED OPEN DATA APPLICATION FOR MOVIES']:
        paragraph(doc,text,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,'Supervisor: '+SUPERVISOR,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,'Team members and student IDs',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    for name,student_id in MEMBERS:
        paragraph(doc,name+' — '+student_id,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,'English report · MovieLOD ontology 3.0.0',align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)
    paragraph(doc,'9 October 2026',align=WD_ALIGN_PARAGRAPH.CENTER,first_line=False)

def markdown():
    lines=['# MovieLOD: course project report (English)','', 'Formatting: Times New Roman 13 pt; 1.5 line spacing; justified body text; A4; left 3 cm, right/top/bottom 2 cm. PDF has 15 pages including cover, contents and references.','']
    for number,page in enumerate(PAGES,1):
        lines += [f'## Page {number}: {page["title"]}','']
        if page.get('kind')=='cover':lines += ['Hanoi University of Science and Technology','Supervisor: '+SUPERVISOR]+[name+' — '+student_id for name,student_id in MEMBERS]+['']
        if page.get('kind')=='toc':
            for index,item in enumerate(PAGES[2:],3):
                lines += [item['title']+' … '+str(index)]
                lines += ['  '+b['text']+' … '+str(index) for b in item['blocks'] if b['type']=='heading']
            lines += ['']
        for block in page['blocks']:
            kind=block['type']
            if kind=='paragraph':lines += [block['text'],'']
            elif kind=='heading':lines += ['### '+block['text'],'']
            elif kind=='code':lines += ['```',block['text'],'```','']
            elif kind=='image':lines += ['!['+block['caption']+']('+str(Path(block['path']).relative_to('docs'))+')','']
            elif kind=='table':
                if block.get('caption'):lines += [block['caption'],'']
                lines += ['| '+' | '.join(block['headers'])+' |','| '+' | '.join(['---']*len(block['headers']))+' |']
                lines += ['| '+' | '.join(str(x) for x in row)+' |' for row in block['rows']]+['']
        lines += ['<!-- page break -->','']
    return '\n'.join(lines)

def main():
    from report_diagrams import main as make_diagrams
    make_diagrams()
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.left_margin=Cm(3);section.right_margin=Cm(2);section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    section.header_distance=Cm(.8);section.footer_distance=Cm(.8);section.different_first_page_header_footer=True
    for style in doc.styles:
        if style.type==1:
            style.font.name=FONT;style.font.size=Pt(13)
            style.paragraph_format.line_spacing=1.5
    normal=doc.styles['Normal'];normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent=Cm(1);normal.paragraph_format.space_after=Pt(6)
    hp=section.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;format_run(hp.add_run('MovieLOD · Semantic Web Course Project'),italic=True)
    fp=section.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=fp.add_run('1');format_run(run)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');field.append(run._r);fp._p.append(field)
    for number,page in enumerate(PAGES,1):
        if page.get('kind')=='cover':cover(doc);continue
        title=heading(doc,page['title'])
        title.paragraph_format.page_break_before=number>1
        if page.get('kind')=='toc':
            title.alignment=WD_ALIGN_PARAGRAPH.CENTER
            contents(doc)
            continue
        for block in page['blocks']:
            kind=block['type']
            if kind=='paragraph':paragraph(doc,block['text'])
            elif kind=='heading':heading(doc,block['text'])
            elif kind=='code':paragraph(doc,block['text'],align=WD_ALIGN_PARAGRAPH.LEFT,first_line=False)
            elif kind=='image':figure(doc,block)
            elif kind=='table':add_table(doc,block)
    destination=ROOT/'docs/Bao_cao.docx';doc.save(destination)
    (ROOT/'docs/Bao_cao.md').write_text(markdown(),encoding='utf-8')
    office='/Applications/LibreOffice.app/Contents/MacOS/soffice'
    with tempfile.TemporaryDirectory(prefix='movie-lod-report-') as work:
        profile=Path(work)/'profile';target=Path(work)/'pdf';target.mkdir()
        subprocess.run([office,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(target),str(destination)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
        shutil.copy2(target/'Bao_cao.pdf',ROOT/'docs/Bao_cao.pdf')
    pdf=PdfReader(ROOT/'docs/Bao_cao.pdf')
    for i,page in enumerate(pdf.pages,1):print(i,page.extract_text()[:110].replace('\n',' | '))
    print('REPORT PAGES:',len(pdf.pages))
    report={'language':'English','pages':len(pdf.pages),'target_pages':15,'font':FONT,'font_size_pt':13,'line_spacing':1.5,
            'alignment':'justified body text','paper':'A4','margins_cm':{'left':3,'right':2,'top':2,'bottom':2},
            'figures':sum(b['type']=='image' for p in PAGES for b in p['blocks']),
            'figure_source':'1 figure retained from supplied slide PDF; 10 author-generated diagrams and verified result panels','source_manifest':'evidence/report_image_sources.json',
            'cover_members':MEMBERS,'supervisor':SUPERVISOR,'institution':'Hanoi University of Science and Technology',
            'toc':'Dedicated page; hierarchical paragraphs, dotted leaders and right-aligned page numbers; no table',
            'docx_sha256':hashlib.sha256(destination.read_bytes()).hexdigest(),
            'pdf_sha256':hashlib.sha256((ROOT/'docs/Bao_cao.pdf').read_bytes()).hexdigest(),'video_modified':False}
    write_json(ROOT/'evidence/report_format_checks.json',report)
    if len(pdf.pages)!=15:raise SystemExit('Report pagination needs adjustment; inspect the page summary.')
    from check_report import main as check_report
    check_report()
    return len(pdf.pages)

if __name__=='__main__':main()
