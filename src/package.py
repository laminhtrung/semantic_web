"""Verify current canonical artifacts and package the project without MP4 or local state."""
import hashlib,json,re,zipfile
from pypdf import PdfReader
from pptx import Presentation
from rdflib import Graph
from rdflib.compare import isomorphic
from common import ROOT,write_json

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    data=Graph().parse(ROOT/'data/processed/movies.ttl');schema=Graph().parse(ROOT/'ontology/movie.ttl')
    assert isomorphic(data,Graph().parse(ROOT/'data/processed/movies.jsonld',format='json-ld'))
    assert isomorphic(data+schema,Graph().parse(ROOT/'ontology/Movie_Knowledge_Graph.owl',format='xml'))
    proof=json.loads((ROOT/'evidence/ontology_design/final_owl_checks.json').read_text())
    assert proof['consistent'] and not proof['unsatisfiable_named_classes'] and proof['sha256']==sha(ROOT/proof['file'])
    assert not (ROOT/'ontology/Movie_Knowledge_Graph_Protege.owl').exists()
    slides=len(Presentation(ROOT/'docs/Slide.pptx').slides)
    assert slides==len(PdfReader(ROOT/'docs/Slide.pdf').pages)==24
    assert len(Presentation(ROOT/'docs/Slide_ngan_13.pptx').slides)==13
    fmt=json.loads((ROOT/'evidence/report_format_checks.json').read_text())
    assert len(PdfReader(ROOT/'docs/Bao_cao.pdf').pages)==15 and fmt['verified']
    assert fmt['font']=='Times New Roman' and fmt['font_size_pt']==13 and fmt['line_spacing']==1.5
    assert fmt['pdf_sha256']==sha(ROOT/'docs/Bao_cao.pdf') and fmt['docx_sha256']==sha(ROOT/'docs/Bao_cao.docx')
    lang=json.loads((ROOT/'evidence/presentation_language_checks.json').read_text())
    assert not lang['unexpected_non_english_text'] and lang['pptx_sha256']==sha(ROOT/'docs/Slide.pptx')
    sources=json.loads((ROOT/'data/raw/snapshots.json').read_text())
    assert all((ROOT/s['path']).is_file() and sha(ROOT/s['path'])==s['sha256'] for s in sources)
    pub=json.loads((ROOT/'evidence/publication.json').read_text());public=json.loads((ROOT/'evidence/publication_checks.json').read_text())
    assert pub['status']=='succeeded' and pub['ontology_version']=='3.0.0' and not pub['local_changes_pending_publication']
    assert public['summary']['publication_complete']
    sync=json.loads((ROOT/'evidence/full_sync_checks.json').read_text());assert sync['passed'] and sync['mp4_files']==0
    match=re.search(r'(\d+) passed',(ROOT/'evidence/tests.txt').read_text());assert match
    browser=json.loads((ROOT/'evidence/browser_checks.json').read_text());assert all(c['passed'] for c in browser)
    queries=json.loads((ROOT/'evidence/public_query_checks.json').read_text())
    assert len(queries['checks'])==31 and all(c['passed'] for c in queries['checks']) and queries['invalid_query_recovery'] and not queries['uncaught_errors']
    assert queries['url'].startswith('https://movie-lod-semantic-web.')
    assert not any(p for p in ROOT.rglob('*') if p.is_file() and p.suffix.lower()=='.mp4' and '.git' not in p.parts)
    summary={'ontology_version':'3.0.0','application_version':'3.0.0','turtle_jsonld_same_graph':True,'owl_matches_data_and_schema':True,'hermit_canonical_owl_verified':True,'single_primary_full_owl':True,'report_pages':15,'slides':24,'short_slides':13,'report_language':'English','report_font':'Times New Roman','report_font_size_pt':13,'report_line_spacing':1.5,'tests_passed':int(match.group(1)),'browser_checks_passed':len(browser),'public_query_checks_passed':31,'source_snapshots_verified':len(sources),'protege_photo_placeholders':11,'official_grade':None,'mp4_count':0,'mp4_removed_by_request':True,'public_publication_complete':True}
    write_json(ROOT/'evidence/deliverables.json',summary)
    excluded={'.venv','.git','__pycache__','.pytest_cache','video_parts','video_parts_v2'}
    def include(p):
        rel=p.relative_to(ROOT)
        return not(set(rel.parts)&excluded) and p.name!='.DS_Store' and p.suffix.lower() not in ['.pyc','.tex','.mp4'] and (p.suffix!='.log' or 'ontology_design' in rel.parts) and not p.name.startswith('site_') and not p.name.endswith('.tar.gz')
    files=sorted(p for p in ROOT.rglob('*') if p.is_file() and include(p))
    write_json(ROOT/'evidence/file_hashes.json',{str(p.relative_to(ROOT)):sha(p) for p in files if p.name!='file_hashes.json'})
    dest=ROOT.parent/'movie_lod_complete.zip'
    with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in files:z.write(p,ROOT.name+'/'+str(p.relative_to(ROOT)))
    with zipfile.ZipFile(dest) as z:assert z.testzip() is None and not any(n.lower().endswith('.mp4') for n in z.namelist())
    print(json.dumps(summary,ensure_ascii=False,indent=2));print('Archive:',dest,'bytes:',dest.stat().st_size)
if __name__=='__main__':main()
