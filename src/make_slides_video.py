"""Create the detailed 24-slide editable PPTX and matching PDF. Video recording has its own script."""
import argparse
import re
import math
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from common import ROOT, write_json
from presentation_content_en import SLIDES, PROTEGE_SHOTS

SCALE=120
INK='152B3A'; TEAL='007F82'; GOLD='E8B65D'; PAPER='F3F6F7'; WHITE='FFFFFF'; MUTED='506776'; LINE='DDE6E9'
FONT='Be Vietnam Pro'
FONTDIR=Path.home()/'Library/Fonts'
class Canvas:
    def __init__(self,prs,i,title,subtitle,dark=False):
        self.dark=dark;self.ps=prs.slides.add_slide(prs.slide_layouts[6])
        bg=INK if dark else PAPER
        self.ps.background.fill.solid();self.ps.background.fill.fore_color.rgb=RGBColor.from_string(bg)
        self.im=Image.new('RGB',(1920,1080),'#'+bg);self.d=ImageDraw.Draw(self.im)
        self.rect(0,0,.13,9,TEAL)
        self.text(.65,.35,13,.35,'MOVIELOD   /   SEMANTIC WEB   /   ONTOLOGY 2.0',12,GOLD if dark else TEAL,True)
        self.text(.65,.95,14.7,.9,title,34,WHITE if dark else INK,True)
        self.text(.67,1.88,14.6,.6,subtitle,17,'B8CDD6' if dark else MUTED)
        self.rect(.65,8.43,14.7,.01,'395361' if dark else LINE)
        self.text(.65,8.58,13,.22,'Data checked: 08 Oct 2026  •  Sources: project code and evidence/',10,'B8CDD6' if dark else MUTED)
        self.text(14.5,8.55,1,.3,f'{i:02d} / {len(SLIDES)}',12,GOLD if dark else TEAL,True)
    def rect(self,x,y,w,h,color,rounded=False):
        shape=self.ps.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color);shape.line.fill.background()
        bbox=(int(x*SCALE),int(y*SCALE),int((x+w)*SCALE),int((y+h)*SCALE))
        if rounded:self.d.rounded_rectangle(bbox,radius=18,fill='#'+color)
        else:self.d.rectangle(bbox,fill='#'+color)
    def text(self,x,y,w,h,txt,size=19,color=INK,bold=False):
        f=ImageFont.truetype(str(FONTDIR/('BeVietnamPro-Bold.ttf' if bold else 'BeVietnamPro-Regular.ttf')),round(size*SCALE/72))
        symbol_font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Unicode.ttf',round(size*SCALE/72))
        def runs(value):
            return [(part,symbol_font if part in ['→','↔','↓','≠'] else f) for part in re.split('([→↔↓≠])',value) if part]
        def measure(value):
            return sum(self.d.textlength(part,font=face) for part,face in runs(value))
        lines=[]
        for paragraph in txt.split('\n'):
            line=''
            for word in paragraph.split():
                candidate=(line+' '+word).strip()
                if measure(candidate)>(w*SCALE-8) and line:lines.append(line);line=word
                else:line=candidate
            lines.append(line)
        leading=size/72*1.28
        if len(lines)*leading>h+.12:raise ValueError(f'Text exceeds box: {txt[:80]} ({len(lines)*leading:.2f}>{h})')
        tb=self.ps.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
        for j,line in enumerate(lines):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.28)
            cursor=x*SCALE
            for part,face in runs(line):
                run=p.add_run();run.text=part
                if face==symbol_font:run.font.name='Arial Unicode MS'
                self.d.text((cursor,(y+j*leading)*SCALE),part,font=face,fill='#'+color)
                cursor+=self.d.textlength(part,font=face)
    def card(self,x,y,w,h,title,body,number=None):
        self.rect(x,y,w,h,WHITE,True)
        self.rect(x,y,.05,h,TEAL)
        if number:self.text(x+.25,y+.2,w-.5,.7,number,30,TEAL,True);ty=y+1
        else:ty=y+.25
        self.text(x+.25,ty,w-.5,.7,title,21,INK,True)
        self.text(x+.25,ty+.78,w-.5,h-(ty-y)-.9,body,17,MUTED)
    def node(self,x,y,w,h,label,color=TEAL):
        self.rect(x,y,w,h,color,True);self.text(x+.18,y+.14,w-.36,h-.2,label,18,WHITE,True)
    def arrow(self,x,y,w,label=''):
        self.text(x,y,w,.5,'→',26,TEAL,True)
        if label:self.text(x-.15,y+.5,max(w,1.5),.6,label,11,MUTED)
    def picture(self,path,x,y,w,h):
        im=Image.open(path).convert('RGB');im.thumbnail((round(w*SCALE),round(h*SCALE)))
        px=x+(w-im.width/SCALE)/2;py=y+(h-im.height/SCALE)/2
        self.im.paste(im,(round(px*SCALE),round(py*SCALE)))
        self.ps.shapes.add_picture(str(path),Inches(px),Inches(py),width=Inches(im.width/SCALE),height=Inches(im.height/SCALE))


def height_for(c,txt,w,size):
    face=ImageFont.truetype(str(FONTDIR/'BeVietnamPro-Regular.ttf'),round(size*SCALE/72))
    n=0
    for para in txt.split('\n'):
        line=''
        for word in para.split():
            candidate=(line+' '+word).strip()
            if c.d.textlength(candidate,font=face)>w*SCALE-8 and line:n+=1;line=word
            else:line=candidate
        n+=1
    return n*size/72*1.28


def paragraph(c,x,y,w,text,size=18,color=INK,bold=False):
    h=height_for(c,text,w,size)+.04
    c.text(x,y,w,h,text,size,color,bold)
    return y+h


def note(c,text,y=7.92):
    c.text(.85,y,14.3,.45,text,13,MUTED)


def table(c,x,y,widths,headers,rows,size=16,row_min=.42,max_bottom=8.12,compact=False):
    total=sum(widths)
    c.rect(x,y,total,.43,TEAL)
    pos=x
    for w,head in zip(widths,headers):c.text(pos+.12,y+.08,w-.24,.33,head,size,WHITE,True);pos+=w
    y+=.43
    for i,row in enumerate(rows):
        padding=.065 if compact else .13
        rh=max(row_min,max(height_for(c,str(value),w-.24,size) for w,value in zip(widths,row))+padding)
        if y+rh>max_bottom+.03:raise ValueError(f'Table exceeds slide: {row}')
        c.rect(x,y,total,rh,WHITE if i%2==0 else 'E7EFF1')
        pos=x
        for w,value in zip(widths,row):c.text(pos+.12,y+padding/2,w-.24,rh-padding,str(value),size,INK);pos+=w
        y+=rh
    return y


def proof_photo(c,key,x=8.55,y=2.82,w=6.65,h=5.02):
    first_shape=len(c.ps.shapes)
    spec=PROTEGE_SHOTS[key];directory=ROOT/'evidence/protege'
    stem=Path(spec['file']).stem
    found=next((directory/(stem+ext) for ext in ['.png','.jpg','.jpeg'] if (directory/(stem+ext)).exists()),None)
    frame=c.ps.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    frame.name='PHOTO_'+key;frame.fill.solid();frame.fill.fore_color.rgb=RGBColor.from_string(WHITE);frame.line.color.rgb=RGBColor.from_string(TEAL)
    c.d.rounded_rectangle((int(x*SCALE),int(y*SCALE),int((x+w)*SCALE),int((y+h)*SCALE)),radius=18,fill='#'+WHITE,outline='#'+TEAL,width=2)
    if found:
        c.text(x+.22,y+.2,w-.44,.4,f'PROTÉGÉ SCREENSHOT · {key}',14,TEAL,True)
        c.picture(found,x+.22,y+.8,w-.44,h-1.3)
        c.text(x+.22,y+h-.35,w-.44,.3,spec['title_en'],12,MUTED)
        return True
    c.text(x+.22,y+.18,w-.44,.38,f'ẢNH PROTÉGÉ CẦN BỔ SUNG · {key}',14,TEAL,True)
    c.rect(x+.22,y+.8,w-.44,1.15,'E7EFF1',True)
    c.text(x+.4,y+1.03,w-.8,.8,spec['title'],22,TEAL,True)
    c.text(x+.25,y+2.11,w-.5,.42,spec['file'],13,MUTED)
    current=y+2.64
    for i,step in enumerate(spec['steps'],1):
        current=paragraph(c,x+.26,current,w-.52,f'{i}. {step}',14,MUTED)+.1
    c.text(x+.25,y+h-.37,w-.5,.31,'Thay khung bằng ảnh thật; chi tiết thêm trong Speaker Notes.',11,TEAL)
    for i,shape in enumerate(list(c.ps.shapes)[first_shape:]):
        shape.name='PHOTO_'+key+'_'+str(i)
    return False


def tree(c,x,y,w,items,row_height=.38,size=18):
    c.rect(x,y,w,len(items)*row_height+.4,WHITE,True)
    for i,(depth,label) in enumerate(items):
        yy=y+.17+i*row_height
        if depth:c.rect(x+.19+(depth-1)*.24,yy+.02,.022,row_height-.04,TEAL)
        c.text(x+.28+depth*.24,yy,w-.56-depth*.24,row_height+.03,label,size,TEAL if depth==0 else INK,depth<=1)


def small_card(c,x,y,w,h,title,body,size=17):
    c.rect(x,y,w,h,WHITE,True);c.rect(x,y,.045,h,TEAL)
    title_height=height_for(c,title,w-.46,20)
    c.text(x+.23,y+.16,w-.46,title_height+.03,title,20,TEAL,True)
    body_y=y+.16+title_height+.13
    c.text(x+.23,body_y,w-.46,h-(body_y-y)-.09,body,size,MUTED)


def code(c,x,y,w,h,value,size=16):
    c.rect(x,y,w,h,INK,True)
    c.text(x+.22,y+.18,w-.44,h-.36,value,size,WHITE)


def technical_notes(s):
    """Keep exhaustive reference tables in notes so the deck remains 24 pages."""
    if s['kind']=='properties':
        from rdflib import Graph,RDF,RDFS,OWL,BNode
        from rdflib.collection import Collection
        from common import EX
        schema=Graph().parse(ROOT/'ontology/movie.ttl')
        def name(term):
            if term is None:return '—'
            if isinstance(term,BNode):
                union=schema.value(term,OWL.unionOf)
                if union:return ' or '.join(name(value) for value in Collection(schema,union))
                return 'class expression'
            value=str(term)
            return ('ex:'+value[len(str(EX)):] if value.startswith(str(EX)) else 'dbo:'+value.rsplit('/',1)[-1] if value.startswith('http://dbpedia.org/ontology/') else value.rsplit('/',1)[-1])
        lines=['Reference: all 23 object properties (not intended to be read aloud in full):','| Property | Domain | Range | Inverse / characteristics |','|:--|:--|:--|:--|']
        for prop in sorted(set(schema.subjects(RDF.type,OWL.ObjectProperty)),key=str):
            inverse=schema.value(prop,OWL.inverseOf)
            if inverse is None:inverse=next(schema.subjects(OWL.inverseOf,prop),None)
            flags=name(inverse)+(' · functional' if (prop,RDF.type,OWL.FunctionalProperty) in schema else '')
            lines.append('| '+name(prop)+' | '+name(schema.value(prop,RDFS.domain))+' | '+name(schema.value(prop,RDFS.range))+' | '+flags+' |')
        return '\n'.join(lines)
    if s['kind']=='endpoint':
        report=json.loads((ROOT/'evidence/review_2026-10-08.json').read_text())
        lines=['Reference: 24 queries; values are result-row counts or ASK booleans:','| File | Asserted | With schema / inference |','|:--|:--|:--|']
        for row in report['queries']:lines.append('| '+row['file']+' | '+str(row['asserted_result']).lower()+' | '+str(row['with_saved_inference_result']).lower()+' |')
        lines.append('Queries 14/15 each return one aggregate row. Queries 17/18 have LIMIT 20; full Actor/Filmmaker totals are 769/89. Query 24 returns one asserted type or three with inference.')
        return '\n'.join(lines)
    return ''


def render_slide(c,s):
    kind=s['kind']
    if kind=='cover':
        c.text(.85,3,8.5,1.6,'From movie data\nto a knowledge graph',35,WHITE,True)
        c.text(.85,5.08,8,.7,'Ontology · RDF · Linked Open Data · SPARQL',21,'B8CDD6')
        c.text(.85,6.3,8,.8,'Full 24-slide presentation\nTeam / members / class: [complete before submission]',17,'B8CDD6')
        for i,(n,label) in enumerate([('30','real films'),('42','named classes'),('1,727','external links')]):
            c.rect(10.05,2.95+i*1.5,5.15,1.28,'24424F',True);c.text(10.32,3.11+i*1.5,2.2,.65,n,30,GOLD,True);c.text(12.63,3.32+i*1.5,2.3,.4,label,17,WHITE)
    elif kind=='requirements':
        table(c,.8,2.95,[1.05,4.25,9.25],['ID','Requirement','Deliverable / evidence'],s['rows'],size=20,row_min=.85)
        note(c,'Required: report ≤15 pages · slides · 3–5-minute video. Running examples: Inception and Nolan.')
    elif kind=='pipeline':
        for i,(n,title,body) in enumerate(s['stages']):c.card(.8+i*3.72,2.95,3.43,3.05,title,body,n)
        c.rect(.8,6.35,14.55,1.22,TEAL,True)
        c.text(1.05,6.53,14,.92,'APPLICATION: Web / Flask + RDFLib / Comunica / Terminal\nEndpoint: movies.ttl.  --reasoned: adds schema + inferred_classes.ttl.',18,WHITE,True)
    elif kind=='metrics':
        for i,(n,label) in enumerate(s['metrics']):
            x=.8+(i%3)*4.96;y=2.96+(i//3)*2.18
            c.rect(x,y,4.66,1.91,WHITE,True);c.text(x+.28,y+.2,4.1,.82,n,38,TEAL,True);c.text(x+.3,y+1.22,4.05,.5,label,21,MUTED)
        note(c,s['footer'])
    elif kind=='ontology_overview':
        for i,(n,label) in enumerate(s['groups']):
            y=2.88+i*.65;c.rect(.8,y,7.38,.53,WHITE,True);c.text(1,y+.1,1.65,.36,n,17,TEAL,True);c.text(2.6,y+.1,5.3,.4,label,14 if i==5 else 16,INK)
        c.text(1,7.04,6.95,.73,'3 DBpedia classes + 39 ex: classes\n14 classes receive types; 42 classes explicitly declared.',16,MUTED)
        proof_photo(c,s['photo'])
    elif kind=='hierarchy':
        tree(c,.8,2.86,7.35,s['tree'],row_height=.4,size=18)
        proof_photo(c,s['photo']);note(c,s['footer'])
    elif kind=='genre_award':
        current=2.92
        for lines in [s['genre'],s['awards']]:
            items=[((len(t)-len(t.lstrip()))//2,t.strip()) for t in lines]
            tree(c,.8,current,7.35,items,row_height=.305,size=16)
            current+=len(lines)*.305+.51
        proof_photo(c,s['photo']);note(c,s['footer'])
    elif kind=='contribution':
        for i,(left,relation,right) in enumerate([('Nolan','hasContribution','Contribution'),('Contribution','contributionTo','Inception · Film'),('Contribution','hasRole','DirectorRole')]):
            yy=2.96+i*1.16;c.node(.85,yy,2.35,.76,left);c.arrow(3.45,yy+.03,1.1,relation);c.node(4.72,yy,3.35,.76,right)
        small_card(c,.85,6.61,7.35,1.25,'Each role has its own record','Director / Actor / Writer / Producer. Nolan has 3 roles in Inception.',16)
        proof_photo(c,s['photo'])
    elif kind=='properties':
        table(c,.8,2.93,[2.2,3.13,2.12],['Property','Domain → range','Inverse'],s['rows'],size=14,row_min=.48,max_bottom=7.8)
        proof_photo(c,s['photo']);note(c,'contributionBy / contributionTo / hasRole: functional. Domain/range can infer types; inverse reverses a relation.')
    elif kind=='owl_rules':
        for i,(title,body) in enumerate(s['cards']):small_card(c,.8+(i%2)*7.48,2.95+(i//2)*2.42,7.18,2.13,title,body,18)
    elif kind=='defined_table':
        reason=json.loads((ROOT/'evidence/ontology_reasoning.json').read_text())
        meanings={'ActingContribution':'Contribution with ActorRole','DirectingContribution':'Contribution with DirectorRole','WritingContribution':'Contribution with WriterRole','ProducingContribution':'Contribution with ProducerRole','Actor':'Person with acting contributions','Filmmaker':'Person with filmmaking contributions','AwardWinner':'Person who receives an award','ActionFilm':'Film with ActionGenre','ComedyFilm':'Film with ComedyGenre','DramaFilm':'Film with DramaGenre','ScienceFictionFilm':'Film with ScienceFictionGenre','AwardWinningFilm':'Film that receives an award','MultiGenreFilm':'Film with at least 2 Genre IRIs','FilmStudio':'Company with at least 3 Film IRIs'}
        rows=[[n,meanings[n],reason['inferred_counts'][n],'OWL RL' if n in reason['rl_defined'] else 'IRI counting'] for n in reason['rl_defined']+reason['cardinality_defined']]
        table(c,.8,2.78,[4.15,6.75,1.35,2.3],['Class','Condition in plain language','Count','Method'],rows,size=15,row_min=.323,max_bottom=7.95,compact=True)
        note(c,s['footer'])
    elif kind=='nolan_chain':
        for i,(title,body) in enumerate(s['steps']):
            y=2.9+i*1.57;small_card(c,.8,y,7.35,1.3,title,body,17)
            if i<2:c.text(4.1,y+1.29,.8,.33,'↓',20,TEAL)
        proof_photo(c,s['photo']);note(c,s['footer'])
    elif kind=='cardinality':
        for i,(title,body) in enumerate(s['cards']):small_card(c,.8,2.9+i*2.12,7.35,1.9,title,body,18)
        c.text(1,7.21,6.93,.61,'Distinct IRIs ≠ provably different OWL individuals.',16,TEAL,True)
        proof_photo(c,s['photo']);note(c,s['footer'])
    elif kind=='collection':
        for i,(title,body) in enumerate(s['cards']):small_card(c,.8,2.9+i*1.61,7.25,1.42,title,body,16)
        c.picture(ROOT/s['image'],8.45,2.85,6.85,5.23)
        note(c,'Right: actual Wikidata metadata and response excerpt. Hashes verify integrity, not real-world truth.')
    elif kind=='rdf':
        code(c,.8,2.88,7.35,1.98,s['code'],16)
        table(c,.8,5.08,[4.05,3.3],['Data property','xsd: type'],s['rows'],size=16,row_min=.35,max_bottom=8)
        proof_photo(c,s['photo'])
    elif kind=='lod':
        for i,(star,title,result) in enumerate(s['stars']):
            y=2.9+i*.86;c.rect(.8,y,14.55,.68,WHITE,True);c.text(1.03,y+.13,1.3,.4,star,21,TEAL,True);c.text(2.6,y+.15,6.4,.39,title,19,INK,True);c.text(9.17,y+.15,5.86,.4,result,17,MUTED)
        note(c,s['footer'])
    elif kind=='resource':
        c.picture(ROOT/s['image'],.8,2.85,7.4,5.1)
        proof_photo(c,s['photo']);note(c,s['footer'])
    elif kind=='query':
        code(c,.8,2.95,6.25,3.3,s['code'],14)
        small_card(c,.8,6.56,6.25,1.2,'2010 · 148 minutes · Nolan','SELECT chooses columns; WHERE matches; OPTIONAL preserves rows.',15)
        c.picture(ROOT/s['image'],7.42,2.84,7.93,5.23)
        note(c,s['footer'])
    elif kind=='roles':
        c.picture(ROOT/s['image'],.8,2.85,9.45,5.2)
        for i,(title,body) in enumerate(s['cards']):small_card(c,10.5,2.92+i*1.7,4.86,1.45,title,body,16)
    elif kind=='endpoint':
        code(c,.8,2.85,7.35,2.73,s['code'],14)
        table(c,.8,5.94,[2.22,5.13],['Mode','Nolan types returned'],s['rows'],size=16,row_min=.68,max_bottom=7.9)
        proof_photo(c,s['photo']);note(c,'--reasoned loads schema + saved types. Endpoint: asserted graph by default. SELECT/ASK: JSON; graph queries: Turtle.')
    elif kind=='validation':
        for i,(n,label) in enumerate(s['checks']):
            y=2.93+i*1.15;c.rect(.8,y,6.5,.94,WHITE,True);c.text(1.02,y+.16,2.15,.57,n if i<3 else '8 / 4',25,TEAL,True);c.text(3.44,y+.23,3.52,.52,label,17,INK)
        c.picture(ROOT/s['image'],7.74,2.86,7.55,4.9)
        note(c,'14 tests / 9 browser checks. HermiT/Pellet verify DL separately; the image shows actual execution logs.')
    elif kind=='score_limits':
        table(c,.8,2.95,[1.15,4.9,1.3],['Req.','Evidence','Score'],s['rows'],size=17,row_min=.65,max_bottom=7.6)
        c.rect(.8,7.18,7.35,.7,TEAL,True);c.text(1.06,7.29,6.8,.45,'Total 10/10 · proposed self-score, not an official grade',17,WHITE,True)
        for i,(title,body) in enumerate(s['cards']):small_card(c,8.55,2.94+i*1.63,6.65,1.55,title,body,16)
    elif kind=='closing':
        photos=[(key, spec) for key,spec in PROTEGE_SHOTS.items()]
        mapping={item['photo']:item['number'] for item in SLIDES if 'photo' in item}
        rows=[[key,str(mapping[key]),spec['title_en']] for key,spec in photos]
        table(c,.8,2.86,[1.1,1.1,6.45],['Image','Slide','Evidence to capture'],rows,size=15,row_min=.345,max_bottom=7.75)
        small_card(c,9.75,2.93,5.43,4.9,'Before submission','1. Complete team / member details.\n\n2. Insert 11 genuine screenshots.\n\n3. Save them in evidence/protege/ for automatic replacement.\n\n4. Rehearse with the speaker script and review the video.',18)
        note(c,s['footer'])
    else:raise ValueError('Unknown slide kind: '+kind)


def create_scripts():
    lines=['---','title: "MovieLOD: lời thuyết trình cho bộ 24 slide"','date: "Bản đầy đủ · 08/10/2026"','---','','## Cách dùng','','Bộ chính có **24 slide**, thuyết trình đầy đủ khoảng **18–22 phút**. Nội dung slide và Speaker Notes bằng tiếng Anh; script riêng giữ lời tiếng Việt tương ứng để tập nói. Hướng dẫn trong ô chờ ảnh vẫn bằng tiếng Việt. Bản ngắn 13 trang vẫn ở Slide_ngan_13.pptx/pdf và Script_thuyet_trinh_ngan_13.md/pdf. Video demo 4:50 là tài liệu riêng; không đọc toàn bộ lời slide vào video.','','Tập theo 3 phần: thành viên A slide 1–10; B slide 11–18; C slide 19–24. Nếu chỉ có 2 người, chia sau slide 14. Các con số lấy từ dữ liệu và evidence hiện tại. Khung ảnh Protégé chưa có ảnh thực; không đọc ghi chú chờ bổ sung như kết quả đã chứng minh.','','## Luồng rút gọn 12–15 phút','','Ưu tiên slide 1–5, 9–10, 12–13, 15–17, 19–23. Các trang cây lớp, công thức cardinality và tra cứu IRI có thể dùng khi trả lời câu hỏi.','','## Mục lục','','| Slide | Nội dung |','|:--|:--|']
    for s in SLIDES:lines.append(f"| {s['number']:02d} | {s['title']} |")
    for s in SLIDES:
        lines.extend([f"\n## Slide {s['number']:02d} — {s['title']}\n",'**Lời nói:**\n',s.get('speech_vi',s['speech']),''])
        if 'photo' in s:
            spec=PROTEGE_SHOTS[s['photo']]
            lines.extend([f"**Ảnh cần bổ sung — {s['photo']}:** `{spec['file']}`; mở `ontology/{spec['owl']}`.\n"]+[f'{i}. {step}' for i,step in enumerate(spec['steps'],1)]+['\n**Cần thấy:** '+spec['expect'],''])
        extra=technical_notes(s)
        if extra:lines.extend(['**Tra cứu khi bảo vệ (không đọc toàn bộ):**\n',extra,''])
        lines.append('**Chuyển trang:** '+('Sau đây nhóm chuyển sang '+SLIDES[s['number']]['title'].lower()+'.' if s['number']<len(SLIDES) else 'Nhóm xin mời thầy cô đặt câu hỏi.'))
    lines.extend(['\n## Câu hỏi bảo vệ ngắn\n','**42 lớp khác gì 30 phim?** 42 là lớp trong schema; 30 là cá thể Film trong graph dữ liệu. Protégé Metrics có thể tính thêm lớp ngoài được tham chiếu.\n','**Vì sao cần Contribution?** Một người có nhiều vai trò trong nhiều phim; mỗi bộ người–phim–vai trò có bản ghi riêng.\n','**Vì sao exact cardinality chưa thay Python?** OWL dùng thế giới mở; Python kiểm tra trường bắt buộc của ứng dụng.\n','**COUNT DISTINCT có phải suy luận DL không?** Đếm tên IRI là quy tắc ứng dụng; OWL không mặc định tên khác nhau chỉ cá thể khác nhau.\n','**Endpoint không có Filmmaker có phải lỗi?** Endpoint mặc định graph gốc. Dùng --reasoned nạp schema và kết quả phân loại.\n','**sameAs khác nguồn thế nào?** sameAs là cùng danh tính; sourceSnapshot là xuất xứ; dbo:Film là tái dùng từ vựng.\n','**SQL có trả được các câu hỏi không?** Có, bằng JOIN/view/quy tắc. Giá trị của bài là IRI, từ vựng chung, liên kết, nguồn và định nghĩa ngữ nghĩa.\n','**Ảnh Protégé có chứng minh reasoner đã chạy không?** Ảnh cây asserted/định nghĩa chỉ chứng minh khai báo. Muốn dùng ảnh inferred, phải ghi rõ reasoner và trạng thái/kết quả thật.\n'])
    (ROOT/'docs/Script_thuyet_trinh.md').write_text('\n'.join(lines),encoding='utf-8')
    lines=['---','title: "MovieLOD: checklist ảnh minh chứng Protégé"','date: "11 vị trí trong bộ 24 slide · 08/10/2026"','---','','## Cách chèn','','Có 11 khung ảnh ghi rõ mã, tên file và thao tác ngay trên slide. Lần tạo này không có ảnh Protégé thật: macOS chặn quyền điều khiển giao diện; không thay bằng ảnh dựng. Ảnh ứng dụng Web trong slide là ảnh chụp thực.','','Cách 1: chèn ảnh vào PowerPoint và che/xóa khung chờ cùng phần hướng dẫn của khung đó. Cách 2: lưu đúng tên ở `evidence/protege/` (PNG/JPG/JPEG cùng stem), chạy `.venv/bin/python src/make_slides_video.py --slides-only`; khung chờ tự thay bằng ảnh. Speaker Notes giữ quy trình chụp.','','Chụp đúng cửa sổ/view, chữ đủ lớn (gợi ý 1440×900 trở lên); không lấy cả desktop có ứng dụng khác. Giữ tên ontology và IRI/thông tin cần chứng minh. Ảnh cây khai báo không được ghi nhãn inferred.','','## Danh sách ảnh\n']
    for s in SLIDES:
        if 'photo' not in s:continue
        key=s['photo'];spec=PROTEGE_SHOTS[key]
        lines.extend([f"### {key} — Slide {s['number']:02d}: {spec['title']}\n",f"**Tên file:** `{spec['file']}`. **Mở:** `ontology/{spec['owl']}`.\n"]+[f'{i}. {step}' for i,step in enumerate(spec['steps'],1)]+['\n**Mục cần thấy:** '+spec['expect'],''])
    lines.extend(['## Kiểm tra trước khi dùng làm minh chứng','','- Đúng ontology 2.0 của bài, không phải file mở rộng cũ.','- Ảnh cá thể dùng Knowledge Graph; ảnh mô hình dùng schema OWL.','- IRI lớp DBpedia, cardinality và datatype đọc được.','- Không nói HermiT/Pellet đã chứng minh DL nếu chỉ chụp công thức hoặc chọn menu.','- Nếu chạy reasoner, giữ tên và trạng thái; ghi cả lỗi nếu có.','','Nguồn hướng dẫn giao diện: [Protégé Views](https://protegeproject.github.io/protege/views/). Ngữ nghĩa cần nhớ: [OWL 2 Primer](https://www.w3.org/TR/owl2-primer/). Nội dung lớp/thuộc tính/cá thể lấy từ OWL của bài.'])
    (ROOT/'docs/Checklist_anh_Protege.md').write_text('\n'.join(lines),encoding='utf-8')
    (ROOT/'evidence/protege/README.md').write_text('# Ảnh Protégé\n\nLưu 11 ảnh theo Checklist_anh_Protege.md/pdf. Đặt đúng stem P00_… tới P10_… để trình tạo slide tự thay khung chờ. Hỗ trợ PNG/JPG/JPEG. Hiện chưa có ảnh thật; không dùng ảnh dựng làm minh chứng.\n',encoding='utf-8')


def create_slides():
    assert len(SLIDES)==24
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    out=ROOT/'docs/slides';out.mkdir(exist_ok=True)
    pages=[];photos=[]
    for s in SLIDES:
        c=Canvas(prs,s['number'],s['title'],s['subtitle'],s['kind']=='cover')
        render_slide(c,s)
        notes=s['speech']
        if 'photo' in s:
            spec=PROTEGE_SHOTS[s['photo']]
            notes+='\n\nẢNH PROTÉGÉ CẦN BỔ SUNG '+s['photo']+' — '+spec['file']+'\nMở ontology/'+spec['owl']+'\n'+'\n'.join(f'{i}. {step}' for i,step in enumerate(spec['steps'],1))+'\nCần thấy: '+spec['expect']
            available=any((ROOT/'evidence/protege'/(Path(spec['file']).stem+ext)).exists() for ext in ['.png','.jpg','.jpeg'])
            photos.append({'id':s['photo'],'slide':s['number'],'file':spec['file'],'image_present':available,'instructions_visible_on_slide':not available})
        c.ps.notes_slide.notes_text_frame.text=notes
        extra=technical_notes(s)
        if extra:c.ps.notes_slide.notes_text_frame.text+='\n\n'+extra
        c.im.save(out/f"{s['number']:02d}.png");pages.append(c.im)
    prs.save(ROOT/'docs/Slide.pptx')
    # Prefer the actual editable PowerPoint layout as a searchable vector PDF.
    office=Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    if office.exists():
        with tempfile.TemporaryDirectory(prefix='movie-lod-slide-export-') as work:
            profile=Path(work)/'profile';target=Path(work)/'pdf';target.mkdir()
            subprocess.run([str(office),'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(target),str(ROOT/'docs/Slide.pptx')],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
            shutil.copy2(target/'Slide.pdf',ROOT/'docs/Slide.pdf')
        import fitz
        pdf=fitz.open(ROOT/'docs/Slide.pdf');assert len(pdf)==len(SLIDES)
        pages=[]
        for i,page in enumerate(pdf):
            pix=page.get_pixmap(matrix=fitz.Matrix(120/72,120/72));pix.save(out/f'{i+1:02d}.png');pages.append(Image.open(out/f'{i+1:02d}.png').convert('RGB'))
        export='LibreOffice vector PDF; searchable text'
    else:
        pages[0].save(ROOT/'docs/Slide.pdf',save_all=True,append_images=pages[1:],resolution=120);export='Raster fallback from matching slide layout'
    contact=Image.new('RGB',(1920,math.ceil(len(pages)/4)*270),'#'+WHITE)
    for i,im in enumerate(pages):contact.paste(im.resize((480,270)),((i%4)*480,(i//4)*270))
    contact.save(out/'contact_sheet.jpg')
    create_scripts()
    write_json(ROOT/'evidence/presentation.json',{'review_date':'2026-10-08','slide_count':len(SLIDES),'font':FONT,'slide_language':'English','placeholder_language':'Vietnamese','speaker_notes_language':'English','rehearsal_script_language':'Vietnamese','editable_text_and_diagrams':True,'pdf_export':export,'speaking_time_minutes':'18–22','protege_photos':photos,'protege_images_present':sum(p['image_present'] for p in photos),'protege_placeholders':sum(not p['image_present'] for p in photos),'real_application_screenshots':['05_inception_current.png','06_nolan_roles.png','09_inception_resource.png'],'short_deck_preserved':'docs/Slide_ngan_13.pptx','technical_self_score':10.0})
    print('Created 24-slide PPTX/PDF, matching speaker script and Protégé checklist.')

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--slides-only',action='store_true',help='Compatibility flag: this generator creates slides, notes and Markdown scripts.');parser.parse_args();create_slides()
