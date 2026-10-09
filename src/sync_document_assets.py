"""Copy canonical documents and create a presentation bundle without MP4."""
import hashlib,json,shutil,zipfile
from pypdf import PdfReader
from pptx import Presentation
from common import ROOT,write_json

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    docs=ROOT/'docs';web=ROOT/'web/dist/docs';web.mkdir(exist_ok=True)
    names=['Slide.pptx','Slide.pdf','Slide_ngan_13.pptx','Slide_ngan_13.pdf','Bao_cao.docx','Bao_cao.pdf','Huong_dan_A_Z.pdf','Huong_dan_thao_tac_chi_tiet.pdf','Mo_ta_ontology.pdf','Ontology_Redesign.pdf','Ket_qua_reasoner.pdf','Script_thuyet_trinh.pdf','Script_thuyet_trinh_ngan_13.pdf','Checklist_anh_Protege.pdf','Kich_ban_video.pdf','Huong_dan_doc_hieu_project.pdf','Huong_dan_doc_hieu_project.docx','DBpedia_OWL_Design.html','DBpedia_OWL_Design.docx','Huong_dan_doc_hieu_project.html']
    for name in names:shutil.copy2(docs/name,web/name)
    images=web/'guide_images';images.mkdir(exist_ok=True)
    for p in (docs/'guide_images').glob('*.png'):shutil.copy2(p,images/p.name)
    shutil.copy2(ROOT/'CHAM_DIEM.pdf',web/'CHAM_DIEM.pdf')
    bundle=docs/'Bo_tai_lieu_thuyet_trinh.zip'
    with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
        for name in names:z.write(docs/name,'docs/'+name)
        for p in sorted(docs.glob('*.md')):z.write(p,'docs/'+p.name)
        for name in ['DBpedia_OWL_Design.zip','Huong_dan_doc_hieu_project.zip','Slide_full.pptx.pdf']:z.write(docs/name,'docs/'+name)
        for folder in ['guide_images','sync_images','report_images']:
            for p in sorted((docs/folder).glob('*')):
                if p.is_file():z.write(p,str(p.relative_to(ROOT)))
        for name in ['Movie_Knowledge_Graph.owl','Movie_Ontology.owl','movie.ttl']:z.write(ROOT/'ontology'/name,'ontology/'+name)
        for p in sorted((ROOT/'queries').glob('*.rq')):z.write(p,str(p.relative_to(ROOT)))
        for name in ['README.md','CHAM_DIEM.md','CHAM_DIEM.pdf','config.json','requirements.txt','requirements_reasoner.txt','Makefile']:z.write(ROOT/name,name)
        for name in ['common.py','query.py']:z.write(ROOT/'src'/name,'src/'+name)
        for name in ['movies.ttl','reasoned.ttl','before_after.trig']:z.write(ROOT/'data/processed'/name,'data/processed/'+name)
        for name in ['final_owl_checks.json','query_results.json','source_mapping.json','checks.json']:
            p=ROOT/'evidence/ontology_design'/name;z.write(p,str(p.relative_to(ROOT)))
    with zipfile.ZipFile(bundle) as z:assert z.testzip() is None and not any(n.lower().endswith('.mp4') for n in z.namelist())
    assert not any(p for p in ROOT.rglob('*') if p.is_file() and p.suffix.lower()=='.mp4' and '.git' not in p.parts)
    for name in names:assert sha(docs/name)==sha(web/name),name
    assert len(Presentation(docs/'Slide.pptx').slides)==len(PdfReader(docs/'Slide.pdf').pages)==24
    assert len(Presentation(docs/'Slide_ngan_13.pptx').slides)==len(PdfReader(docs/'Slide_ngan_13.pdf').pages)==13
    assert len(PdfReader(docs/'Bao_cao.pdf').pages)==15
    fmt=json.loads((ROOT/'evidence/report_format_checks.json').read_text());assert fmt['verified']
    language=json.loads((ROOT/'evidence/presentation_language_checks.json').read_text());assert language['pptx_sha256']==sha(docs/'Slide.pptx') and language['pdf_sha256']==sha(docs/'Slide.pdf')
    owl=json.loads((ROOT/'evidence/ontology_design/final_owl_checks.json').read_text());assert owl['sha256']==sha(ROOT/'ontology/Movie_Knowledge_Graph.owl')
    report={'passed':True,'ontology_version':'3.0.0','application_version':'3.0.0','slide_pages':24,'short_slide_pages':13,'report_pages':15,'slide_language':'English','photo_instructions_language':'Vietnamese','script_language':'Vietnamese','protege_placeholders':11,'report_format_preserved':True,'local_web_document_copies_current':True,'current_mp4_count':0,'mp4_removed_by_request':True,'document_assets':{name:sha(docs/name) for name in names},'bundle_sha256':sha(bundle),'retained_source_pdf':'docs/Slide_full.pptx.pdf; provenance source, not the current deck'}
    write_json(ROOT/'evidence/document_sync_checks.json',report);print(json.dumps({k:v for k,v in report.items() if k!='document_assets'},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
