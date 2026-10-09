"""Build clean Word tables and paragraphs from Pandoc's Markdown AST."""
import json,subprocess
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

def plain(inlines):
    text=''
    for item in inlines:
        t=item['t'];c=item.get('c')
        if t=='Str':text+=c
        elif t in ['Space','SoftBreak','LineBreak']:text+=' '
        elif t=='Code':text+=c[1]
        elif t in ['Strong','Emph','SmallCaps','Strikeout']:text+=plain(c)
        elif t in ['Link','Image','Span']:text+=plain(c[1])
        elif t=='Quoted':text+='“'+plain(c[1])+'”'
    return text

def inline(p,items,bold=False,italic=False):
    for item in items:
        t=item['t'];c=item.get('c')
        if t in ['Strong','Emph']:inline(p,c,bold or t=='Strong',italic or t=='Emph');continue
        if t=='Link':
            element=OxmlElement('w:hyperlink');element.set(qn('r:id'),p.part.relate_to(c[2][0],RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'006E88');prop.append(color);run.append(prop)
            text=OxmlElement('w:t');text.text=plain(c[1]);run.append(text);element.append(run);p._p.append(element);continue
        if t=='Str':text=c
        elif t in ['Space','SoftBreak']:text=' '
        elif t=='LineBreak':text='\n'
        elif t=='Code':text=c[1]
        elif t=='Quoted':text='“'+plain(c[1])+'”'
        elif t=='Span':inline(p,c[1],bold,italic);continue
        else:continue
        run=p.add_run(text);run.bold=bold;run.italic=italic
        if t=='Code':run.font.name='Menlo';run.font.size=Pt(10)

def create(md,reference):
    ast=json.loads(subprocess.check_output(['pandoc',str(md),'--to=json']))
    d=Document(reference)
    meta=ast.get('meta',{})
    for key,style in [('title','Title'),('subtitle','Subtitle'),('date','Normal')]:
        item=meta.get(key)
        if item:d.add_paragraph(plain(item['c']),style)
    index=0
    def blocks(items,container=None):
        nonlocal index
        target=container or d
        for block in items:
            t=block['t'];c=block.get('c')
            if t=='Header':inline(target.add_paragraph(style='Heading '+str(c[0])),c[2])
            elif t in ['Para','Plain']:
                if len(c)==1 and c[0]['t']=='Image':
                    img=c[0]['c'];p=target.add_paragraph();p.add_run().add_picture(str(md.parent/img[2][0]),width=Cm(15));p=target.add_paragraph(plain(img[1]));p.runs[0].italic=True
                else:inline(target.add_paragraph(),c)
            elif t=='CodeBlock':
                p=target.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6)
                r=p.add_run(c[1]);r.font.name='Menlo';r.font.size=Pt(10)
            elif t in ['BulletList','OrderedList']:
                content=c if t=='BulletList' else c[1]
                for number,item in enumerate(content,1):
                    for j,entry in enumerate(item):
                        if entry['t'] in ['Plain','Para']:
                            p=target.add_paragraph();p.add_run(('• ' if t=='BulletList' else str(number)+'. ') if j==0 else '');inline(p,entry['c'])
                        else:blocks([entry],target)
            elif t=='Table':
                rows=list(c[3][1])
                for body in c[4]:rows+=body[2]+body[3]
                rows+=c[5][1]
                table=d.add_table(rows=0,cols=len(c[2]));table.style='Table Grid'
                for row in rows:
                    cells=table.add_row().cells
                    for cell,value in zip(cells,row[1]):
                        entries=value[4]
                        first=True
                        for entry in entries:
                            if entry['t'] in ['Para','Plain']:
                                p=cell.paragraphs[0] if first else cell.add_paragraph();first=False;inline(p,entry['c'])
                            else:blocks([entry],cell)
                for cell in table.rows[0].cells:
                    for p in cell.paragraphs:
                        for r in p.runs:r.bold=True
                d.add_paragraph();index+=1
            elif t=='Div':blocks(c[1],target)
            elif t=='Figure':blocks(c[2],target)
            else:raise ValueError('Unsupported Markdown block: '+t)
    blocks(ast['blocks'])
    return d
