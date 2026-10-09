"""Render a Vietnamese learning guide and verify its examples against real graphs."""
import hashlib,json,re,subprocess,tempfile
from pathlib import Path
from rdflib import Graph,Dataset,RDF,OWL,URIRef
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from common import ROOT,EX,RES,DBO,write_json

NAME='Huong_dan_doc_hieu_project'
OUT=ROOT/'evidence/reading_guide'

def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def diagrams():
    folder=ROOT/'docs/guide_images';folder.mkdir(exist_ok=True)
    graphs={
    '01_pipeline':'''digraph G {rankdir=LR; node [shape=box,style="rounded,filled",fillcolor="#E9F4F5",fontname="Times New Roman",fontsize=16,color="#247A86"]; edge [color="#247A86",penwidth=2];
    source [label="Crawled claims\nWikidata + source snapshots"];
    rdf [label="Asserted RDF\nFilms, people, credits"];
    owl [label="OWL schema\nDBpedia reuse + extensions",fillcolor="#FEF0D8"];
    reason [label="HermiT / OWL RL\nEntailed types / relations"];
    query [label="SPARQL\nCompetency questions"];
    source->rdf; rdf->reason; owl->reason; reason->query; rdf->query [label="direct queries",fontname="Times New Roman",fontsize=12];
    }''',
    '02_contribution':'''digraph G {rankdir=LR;node [shape=box,style="rounded,filled",fillcolor="#E9F4F5",fontname="Times New Roman",fontsize=16,color="#247A86"];edge [color="#247A86",fontname="Times New Roman",fontsize=12];
    person [label="Christopher Nolan\ndbo:Person"];
    film [label="Inception\ndbo:Film"];
    d [label="Directing credit"];
    w [label="Writing credit"];
    p [label="Producing credit"];
    dr [label="ex:DirectorRole",fillcolor="#FEF0D8"];
    wr [label="ex:WriterRole",fillcolor="#FEF0D8"];
    pr [label="ex:ProducerRole",fillcolor="#FEF0D8"];
    person->d [label="hasContribution"];person->w;person->p;
    d->film [label="contributionTo"];w->film;p->film;
    d->dr [label="hasRole"];w->wr [label="hasRole"];p->pr [label="hasRole"];
    {rank=same;d;w;p;} {rank=same;dr;wr;pr;}
    }''',
    '03_reasoning':'''digraph G {rankdir=TB;node [shape=box,style="rounded,filled",fillcolor="#E9F4F5",fontname="Times New Roman",fontsize=16,color="#247A86"];edge [color="#247A86",fontname="Times New Roman",fontsize=12];
    role [label="Facts: Nolan has directing + writing credits\nCredit roles: DirectorRole / WriterRole",fillcolor="#FEF0D8"];
    d [label="DirectingContribution\nhasRole value DirectorRole"];
    w [label="WritingContribution\nhasRole value WriterRole"];
    md [label="dbo:MovieDirector\ncredit existential restriction"];
    sw [label="dbo:ScreenWriter\ncredit existential restriction"];
    wd [label="ex:WriterDirector\nMovieDirector AND ScreenWriter",fillcolor="#D7ECE0"];
    fm [label="ex:Filmmaker\ndirecting OR writing OR producing",fillcolor="#D7ECE0"];
    role->d;role->w;d->md;w->sw;md->wd;sw->wd;d->fm;w->fm;
    }'''}
    for name,dot in graphs.items():
        file=folder/(name+'.dot');file.write_text(dot)
        subprocess.run(['dot','-Tpng','-Gdpi=180',str(file),'-o',str(folder/(name+'.png'))],check=True)
    return [folder/(n+'.png') for n in graphs]

def examples():
    folder=ROOT/'evidence/ontology_design';a=Graph().parse(folder/'asserted.ttl');g=Graph().parse(folder/'schema.ttl')+a+Graph().parse(folder/'inferred.ttl')
    prefix=f'PREFIX dbo: <{DBO}>\nPREFIX ex: <{EX}>\nPREFIX res: <{RES}>\nPREFIX rdfs: <http://www.w3.org/2000/01/rdf-schema#>\n'
    local='FILTER(STRSTARTS(STR(?film), "'+str(RES)+'"))'
    direct=prefix+'SELECT DISTINCT ?film ?name WHERE { ?film a dbo:Film ; dbo:director res:person-Q25191 ; rdfs:label ?name . '+local+' } ORDER BY ?name'
    writer=prefix+'SELECT DISTINCT ?person ?name WHERE { ?person a ex:WriterDirector ; rdfs:label ?name . FILTER(STRSTARTS(STR(?person), "'+str(RES)+'")) } ORDER BY ?name'
    ask=prefix+'ASK { res:person-Q25191 a ex:ThreeCreditContributor }'
    chain=prefix+'SELECT DISTINCT ?film WHERE { res:person-Q25191 ex:contributedTo ?film . '+local+' } ORDER BY ?film'
    results={}
    for name,q in [('nolan_films',direct),('writer_directors',writer),('contributed_to',chain)]:results[name]={'asserted':len(list(a.query(q))),'reasoned':len(list(g.query(q)))}
    results['nolan_min3']={'asserted':bool(a.query(ask)),'reasoned':bool(g.query(ask))}
    assert results['nolan_films']['asserted']==8
    assert results['writer_directors']=={'asserted':0,'reasoned':10}
    assert results['contributed_to']=={'asserted':0,'reasoned':8}
    assert results['nolan_min3']=={'asserted':False,'reasoned':True}
    ds=Dataset().parse(folder/'before_after.trig',format='trig');q=(ROOT/'queries/design/27.rq').read_text();results['query27_rows']=len(list(ds.query(q)))
    expected=json.loads((folder/'query_results.json').read_text());assert results['query27_rows']==expected[-1]['reasoned_rows']
    owl=Graph().parse(ROOT/'ontology/Movie_Knowledge_Graph.owl');assert len(owl)==19025
    assert len({c for c in owl.subjects(RDF.type,OWL.Class) if isinstance(c,URIRef)})==37
    assert len(set(owl.subjects(RDF.type,EX.ContributionRole)))==4
    f=RES['film-Q25188'];n=RES['person-Q25191']
    assert float(owl.value(f,DBO.runtime))==8880 and int(owl.value(f,EX.releaseYear))==2010
    credits=[c for c in owl.objects(n,EX.hasContribution) if owl.value(c,EX.contributionTo)==f]
    assert {owl.value(c,EX.hasRole) for c in credits}=={EX.DirectorRole,EX.WriterRole,EX.ProducerRole}
    assert len(set(owl.objects(n,EX.hasContribution)))==22
    return results

def reference():
    d=Document();normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(13)
    normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for n,size in [('Title',22),('Subtitle',14),('Heading 1',16),('Heading 2',14),('Heading 3',13)]:
        st=d.styles[n];st.font.name='Times New Roman';st.font.size=Pt(size);st.paragraph_format.line_spacing=1.5
        if n.startswith('Heading'):st.paragraph_format.keep_with_next=True
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(3);sec.right_margin=Cm(2);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
        h=sec.header.paragraphs[0];h.text='MOVIELOD • HƯỚNG DẪN ĐỌC HIỂU • ONTOLOGY 3.0.0';h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        for r in h.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    file=OUT/'reference.docx';d.save(file);return file

def main():
    OUT.mkdir(parents=True,exist_ok=True);doc=ROOT/'docs'/(NAME+'.md')
    primary=ROOT/'ontology/Movie_Knowledge_Graph.owl';inputs_before=digest(primary)
    images=diagrams();results=examples()
    css='''body{font-family:"Times New Roman",serif;font-size:19px;line-height:1.65;color:#203442;max-width:1100px;margin:24px auto;padding:24px}h1,h2{color:#075a66;line-height:1.3}h1{margin-top:40px;border-bottom:2px solid #d7e9eb;padding-bottom:8px}a{color:#006e88}table{display:block;overflow-x:auto;width:100%;border-collapse:collapse;font-size:17px}th,td{padding:10px;border:1px solid #c6d7dc;text-align:left;vertical-align:top}th{background:#eaf4f5}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#edf3f5;padding:18px;line-height:1.5;font-size:14px}code{font-family:Menlo,monospace;font-size:.8em;overflow-wrap:anywhere}img{max-width:100%;height:auto}nav{background:#edf7f8;padding:20px}p{overflow-wrap:anywhere}'''
    header=OUT/'header.html';header.write_text('<style>'+css+'</style>')
    subprocess.run(['pandoc',str(doc),'--standalone','--toc','--toc-depth=2','--include-in-header='+str(header),'-o',str(doc.with_suffix('.html'))],check=True,cwd=ROOT/'docs')
    from reading_guide_word import create
    word=create(doc,reference())
    # Remove theme font overrides so headings actually use Times New Roman.
    for style in word.styles:
        if style.type in [1,2]:
            style.font.name='Times New Roman'
            fonts=style._element.find('.//'+qn('w:rFonts'))
            if fonts is not None:
                for key in list(fonts.attrib):
                    if 'Theme' in key or 'theme' in key:del fonts.attrib[key]
    for p in word.paragraphs:
        p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0);p.paragraph_format.first_line_indent=Cm(0)
        if not p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
        if p.style.name in ['Source Code','SourceCode','Verbatim']:
            p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(2)
            for r in p.runs:r.font.name='Menlo';r.font.size=Pt(10)
    widths=[[3,5,8],[3,5,8],[4,12],[2.5,5,8.5],[3,5,8],[6,2,8],[4,7,5],[7,9],[7,2,7],[3.5,6.5,6],[6,10],[8,8]]
    for index,table in enumerate(word.tables):
        table.style='Table Grid';table.autofit=False
        selected=widths[index]
        assert len(selected)==len(table.columns)
        for col,w in zip(table.columns,selected):col.width=Cm(w)
        table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
        table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),'9071')
        for row in table.rows:
            for cell,w in zip(row.cells,selected):cell.width=Cm(w)
            trPr=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name='Times New Roman';run.font.size=Pt(12)
    headings=[p for p in word.paragraphs if p.style.name=='Heading 1']
    first=headings[0]
    for element in list(word._element.body):
        if element is first._p:break
        if element.tag==qn('w:sdt'):word._element.body.remove(element)
    def before(text):
        node=OxmlElement('w:p');first._p.addprevious(node);p=Paragraph(node,first._parent);p.add_run(text);return p
    toc_title=before('MỤC LỤC');toc_title.paragraph_format.page_break_before=True;toc_title.alignment=WD_ALIGN_PARAGRAPH.CENTER;toc_title.runs[0].bold=True
    toc=[]
    for heading in headings:
        p=before(heading.text+'\t');p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS);toc.append(p)
    first.paragraph_format.page_break_before=True
    word.save(doc.with_suffix('.docx'))
    with tempfile.TemporaryDirectory(prefix='movie-reading-guide-') as temp:
        profile=Path(temp)/'profile';pdfout=Path(temp)/'pdf';pdfout.mkdir()
        subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(pdfout),str(doc.with_suffix('.docx'))],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,timeout=120)
        initial=PdfReader(pdfout/(NAME+'.pdf'))
        texts=[' '.join((page.extract_text() or '').split()) for page in initial.pages]
        toc_pages={}
        for heading,p in zip(headings,toc):
            normalized=' '.join(heading.text.split());number=next((i+1 for i,t in enumerate(texts) if i>=2 and normalized in t),None)
            assert number is not None,heading.text
            p.add_run(str(number));toc_pages[heading.text]=number
        word.save(doc.with_suffix('.docx'))
        subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(pdfout),str(doc.with_suffix('.docx'))],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,timeout=120)
        doc.with_suffix('.pdf').write_bytes((pdfout/(NAME+'.pdf')).read_bytes())
    pdf=PdfReader(doc.with_suffix('.pdf'));text='\n'.join(p.extract_text() or '' for p in pdf.pages)
    assert len(pdf.pages)>10 and 'WriterDirector' in text and 'Contribution' in text
    assert len(Document(doc.with_suffix('.docx')).inline_shapes)==3
    assert inputs_before==digest(primary)
    final=json.loads((ROOT/'evidence/ontology_design/final_owl_checks.json').read_text());assert final['sha256']==inputs_before
    write_json(OUT/'checks.json',{'passed':True,'language':'Vietnamese','ontology_version':'3.0.0','owl_sha256':inputs_before,'canonical_owl_unchanged':True,'mp4_included':False,'examples':results,'pdf_pages':len(pdf.pages),'toc_pages':toc_pages,'authored_diagrams':3,'diagram_note':'Conceptual illustrations derived from real facts and axioms, not app/Protege screenshots','artifacts':{str(p.relative_to(ROOT)):digest(p) for p in [doc,doc.with_suffix('.html'),doc.with_suffix('.docx'),doc.with_suffix('.pdf')]+images}})
    print(json.dumps({'pages':len(pdf.pages),'examples':results,'file':str(doc.with_suffix('.pdf'))},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
